from telegram.ext import ApplicationBuilder, CommandHandler
from openpyxl import load_workbook
from dotenv import load_dotenv
from datetime import datetime
import os

load_dotenv("config.env")

TOKEN = os.getenv("BOT_TOKEN")
OWNER = os.getenv("OWNER_NAME")
EXCEL_FILE = os.getenv("EXCEL_FILE")

print("TOKEN:", TOKEN)


def save_to_excel(category, amount):

    wb = load_workbook(EXCEL_FILE)
    ws = wb["Pengeluaran Bulanan"]

    row = 2

    while ws[f"A{row}"].value is not None:
        row += 1

    ws[f"A{row}"] = datetime.now().strftime("%d-%m-%Y")
    ws[f"B{row}"] = OWNER
    ws[f"C{row}"] = category
    ws[f"D{row}"] = 1
    ws[f"E{row}"] = amount * 1000

    wb.save(EXCEL_FILE)

async def add(update, context, category):

    print("ARGS:", context.args)

    try:
        amount = int(context.args[0])

        save_to_excel(category, amount)

        await update.message.reply_text(
            f"✅ {category} Rp {amount}.000 tersimpan"
        )

    except Exception as e:

        print("ERROR:", e)

        await update.message.reply_text(
            "Contoh: /kopi 25"
        )

async def kopi(update, context):
    await add(update, context, "Kopi")


async def rokok(update, context):
    await add(update, context, "Rokok")


async def makan(update, context):
    await add(update, context, "Makan")


async def bensin(update, context):
    await add(update, context, "Bensin")


app = ApplicationBuilder().token(TOKEN).build()

app.add_handler(CommandHandler("kopi", kopi))
app.add_handler(CommandHandler("rokok", rokok))
app.add_handler(CommandHandler("makan", makan))
app.add_handler(CommandHandler("bensin", bensin))

print("Bot running...")

app.run_polling()